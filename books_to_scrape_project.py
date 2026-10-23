# PyCharm project script: Books to Scrape - catalogue analysis and price prediction
# Uses "# %%" cell markers, so PyCharm can run it cell by cell.
# Run from the project folder so the data/, figures/ and report/ folders are created next to this file.

# %% [markdown]
# # Books to Scrape: Catalogue Analysis and Book Price Prediction
#
# **Role:** Junior Data Scientist, Publishing Analytics Team
# **Name / Student number:** Ntokozo 402600609
# **Data source:** https://books.toscrape.com (scraped with BeautifulSoup)
#
# ## Project overview
# The analytics team wants to know what influences book prices and ratings in an online catalogue, and whether price can be predicted from public book information. This notebook covers the whole pipeline:
#
# 1. **Data acquisition**: scrape the whole catalogue (1 000 books, 50 genres) with `requests` + `BeautifulSoup`
# 2. **Data cleaning**: turn the messy scraped text into a clean analytical dataset
# 3. **Visualisation and exploration**: price, rating, genre and stock patterns
# 4. **Predictive modelling**: a Linear Regression that estimates price
# 5. **Reporting**: a short consulting-style report
#
# **Important context:** the website itself warns that *"prices and ratings here were randomly assigned and have no real meaning"*. So I do not expect strong relationships, and I treat this as a test of whether my workflow is sound, not a claim about the real book market. I keep that in mind when interpreting results.
#
# **How to run:** `pip install requests beautifulsoup4 pandas numpy matplotlib seaborn scikit-learn scipy python-docx`, then *Run All*. A full scrape takes roughly 5 to 8 minutes because I add a small delay between requests to be polite to the server. To skip scraping on a re-run, set `RESCRAPE = False` (the saved raw CSV is reused).

# %% [markdown]
# ## 0. Setup

# %%
import os, re, time, warnings
import requests
from bs4 import BeautifulSoup
from urllib.parse import urljoin

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from scipy.stats import spearmanr, kruskal, skew

from sklearn.model_selection import train_test_split, cross_val_score, KFold
from sklearn.linear_model import LinearRegression
from sklearn.dummy import DummyRegressor
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.feature_extraction.text import CountVectorizer, ENGLISH_STOP_WORDS

warnings.filterwarnings("ignore")
sns.set_theme(style="whitegrid")
pd.set_option("display.max_columns", 30)

os.makedirs("data", exist_ok=True)
os.makedirs("figures", exist_ok=True)
os.makedirs("report", exist_ok=True)

# %% [markdown]
# ## 1. Data Acquisition (Web Scraping)
#
# ### 1.1 Plan
# I inspected the site's HTML first. The data is spread over two kinds of page:
#
# | Page type | What it gives me |
# |---|---|
# | **Listing pages** (`catalogue/page-1.html` ... `page-50.html`, 20 books each) | title (full title is in the link's `title` attribute), price, star rating (stored as a CSS class such as `star-rating Three`), short availability text, link to the product page |
# | **Product pages** (one per book) | genre/category (breadcrumb), exact stock quantity ("In stock (22 available)"), product description, UPC, tax and review fields |
#
# The listing pages alone do not give the category, quantity or description, so I scrape listings first and then visit every product page.
#
# I scrape the **entire catalogue** (not only one genre) so the genre comparisons are possible. The rules I follow: set a User-Agent, wait a short time between requests, retry failed requests, and never let one bad page stop the whole run.

# %%
# ---------- configuration ----------
BASE_URL   = "https://books.toscrape.com/"
START_URL  = urljoin(BASE_URL, "catalogue/page-1.html")
DELAY      = 0.2        # seconds between requests (be polite)
MAX_BOOKS  = None       # None = whole catalogue; fir a quick check set MAX_BOOKS to 200
RESCRAPE   = True       # False = reuse data/books_raw.csv if it already exists
RAW_PATH   = "data/books_raw.csv"
CLEAN_PATH = "data/books_clean.csv"

session = requests.Session()
session.headers.update({"User-Agent": "Mozilla/5.0 (student data science project; BeautifulSoup)"})

# %% [markdown]
# ### 1.2 Helper functions

# %%
def get_soup(url, retries=3):
    # Download a page and return a BeautifulSoup object (None if it keeps failing)
    for attempt in range(1, retries + 1):
        try:
            response = session.get(url, timeout=15)
            response.raise_for_status()
            response.encoding = "utf-8"       # stops "£" turning into "Â£"
            return BeautifulSoup(response.text, "html.parser")
        except requests.RequestException as err:
            print(f"  attempt {attempt}/{retries} failed for {url}: {err}")
            time.sleep(1.5 * attempt)
    return None


def text_or_none(tag):
    return tag.get_text(strip=True) if tag is not None else None


def parse_listing_page(soup, page_url):
    # Return (list of book dicts, url of next page or None) for one listing page
    rows = []
    for pod in soup.select("article.product_pod"):
        link = pod.select_one("h3 a")
        rating_tag = pod.select_one("p.star-rating")
        rating_classes = [c for c in rating_tag.get("class", []) if c != "star-rating"] if rating_tag else []
        rows.append({
            "title": link.get("title") if link else None,
            "product_url": urljoin(page_url, link["href"]) if link else None,
            "price_raw": text_or_none(pod.select_one("p.price_color")),
            "rating_word": rating_classes[0] if rating_classes else None,
            "availability_listing": text_or_none(pod.select_one("p.availability")),
        })
    next_link = soup.select_one("li.next a")
    next_url = urljoin(page_url, next_link["href"]) if next_link else None
    return rows, next_url


def parse_product_page(soup):
    # Pull category, stock detail, description and table fields from a product page
    info = {}
    for tr in soup.select("table.table-striped tr"):
        if tr.th is not None and tr.td is not None:
            info[tr.th.get_text(strip=True)] = tr.td.get_text(strip=True)

    crumbs = soup.select("ul.breadcrumb li a")          # Home > Books > Category
    category = crumbs[2].get_text(strip=True) if len(crumbs) >= 3 else None

    description = None
    header = soup.select_one("#product_description")     # description <p> follows this header
    if header is not None:
        para = header.find_next_sibling("p")
        description = text_or_none(para)

    return {
        "upc": info.get("UPC"),
        "category": category,
        "description": description,
        "availability_detail": info.get("Availability"),
        "price_excl_tax_raw": info.get("Price (excl. tax)"),
        "price_incl_tax_raw": info.get("Price (incl. tax)"),
        "tax_raw": info.get("Tax"),
        "num_reviews_raw": info.get("Number of reviews"),
    }

# %% [markdown]
# ### 1.3 Scrape the catalogue

# %%
def scrape_catalogue(max_books=None):
    # Step 1: walk through every listing page (follow the 'next' button)
    listing, url, page = [], START_URL, 1
    while url:
        soup = get_soup(url)
        if soup is None:
            print(f"Giving up on listing page {page}; continuing with what we have.")
            break
        rows, url = parse_listing_page(soup, url)
        listing.extend(rows)
        print(f"Listing page {page}: {len(rows)} books (total {len(listing)})")
        page += 1
        time.sleep(DELAY)
        if max_books and len(listing) >= max_books:
            listing = listing[:max_books]
            break

    # Step 2: visit every product page for the extra fields
    records = []
    for i, row in enumerate(listing, start=1):
        detail = {}
        soup = get_soup(row["product_url"]) if row["product_url"] else None
        if soup is not None:
            detail = parse_product_page(soup)
        records.append({**row, **detail})
        if i % 100 == 0 or i == len(listing):
            print(f"Product pages scraped: {i}/{len(listing)}")
        time.sleep(DELAY)
    return pd.DataFrame(records)


if (not RESCRAPE) and os.path.exists(RAW_PATH):
    raw = pd.read_csv(RAW_PATH)
    print(f"Loaded saved raw data from {RAW_PATH}")
else:
    raw = scrape_catalogue(MAX_BOOKS)
    raw.to_csv(RAW_PATH, index=False)
    print(f"Saved raw data to {RAW_PATH}")

print("Raw shape:", raw.shape)
raw.head()

# %% [markdown]
# ### 1.4 Quick check of the raw data
# Before cleaning, I check that the size requirement is met (at least 200 books) and that every required field (title, price, star rating, stock) was captured.

# %%
required = ["title", "price_raw", "rating_word", "availability_detail"]
print(f"Books scraped: {len(raw)}  (requirement: at least 200)")
print("Requirement met:", len(raw) >= 200)
print("\nMissing values per column:")
print(raw.isna().sum())
print("\nRating words found:", raw["rating_word"].value_counts(dropna=False).to_dict())
print("Availability examples:", raw["availability_detail"].dropna().unique()[:5])

# %% [markdown]
# ## 2. Data Cleaning and Preparation
#
# The scraped data is all text (`"£51.77"`, `"Three"`, `"In stock (22 available)"`). For each problem I note what I found, what I did, and why.
#
# | Problem | Decision | Justification |
# |---|---|---|
# | Duplicate rows | drop on `product_url` | one row per book |
# | Price is text with a currency symbol (and sometimes a stray `Â`) | strip everything except digits and `.`, convert to float | needed for maths; the stray character comes from a text-encoding mismatch |
# | Rating is a word | map One..Five to 1..5 | star ratings are ordered, so integers keep that order |
# | Stock is a sentence | split into `in_stock` (0/1) and `quantity` (number) | model-ready fields; out of stock gets quantity 0 |
# | Missing description | keep the book, fill with empty text, add `has_description` flag | dropping would waste valid price/rating data; I do not invent text |
# | Missing price or rating | drop those rows | price is the target variable, so it cannot be guessed |
# | Description ends with "...more", has odd characters | remove | scraping artefacts, not real content |
# | Many tiny categories | group categories with fewer than 10 books into "Other" | rare dummy variables cause unstable coefficients |
# | Constant columns (product type, tax, review count) | drop | no information, cannot explain price |

# %%
df = raw.copy()
n_start = len(df)

# 1. duplicates
df = df.drop_duplicates(subset="product_url").reset_index(drop=True)
print("Duplicates removed:", n_start - len(df))

# 2. price -> float (keep digits and the decimal point only)
def to_number(series):
    return pd.to_numeric(series.astype(str).str.replace(r"[^0-9.]", "", regex=True), errors="coerce")

df["price"] = to_number(df["price_raw"])
df["price"] = df["price"].fillna(to_number(df["price_incl_tax_raw"]))    # backup source

# 3. rating word -> number
rating_map = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
df["rating"] = df["rating_word"].map(rating_map)

# 4. stock -> in_stock flag + quantity
stock_text = df["availability_detail"].fillna(df["availability_listing"]).fillna("")
df["in_stock"] = stock_text.str.contains("in stock", case=False).astype(int)
df["quantity"] = pd.to_numeric(stock_text.str.extract(r"\((\d+)\s+available\)")[0], errors="coerce")
df.loc[df["in_stock"] == 0, "quantity"] = 0
median_qty = df.loc[df["in_stock"] == 1, "quantity"].median()
df["quantity"] = df["quantity"].fillna(median_qty)          # in stock but number missing

# 5. rows we cannot use (target or key predictor missing)
before = len(df)
df = df.dropna(subset=["price", "rating"]).reset_index(drop=True)
df["rating"] = df["rating"].astype(int)
print("Rows dropped for missing price/rating:", before - len(df))

# 6. category
df["category"] = df["category"].fillna("Unknown").str.strip()

# 7. description
df["description"] = df["description"].fillna("").str.replace("Â", "", regex=False)
df["description"] = df["description"].str.replace(r"\s*\.\.\.more\s*$", "", regex=True).str.strip()
df["has_description"] = (df["description"].str.len() > 0).astype(int)
print("Books with no description:", int((df["has_description"] == 0).sum()))

# %% [markdown]
# ### 2.2 Text preparation and feature engineering
# The brief mentions descriptions and tags. The site has **no tag field**, so I use two things as "tags": the **category** (the genre label) and **keywords** extracted from descriptions. The text is cleaned (lower case, letters only, stop-words removed) and I create simple numeric features from the title and description. Each feature is something a retailer would know about a listing.

# %%
STOP = set(ENGLISH_STOP_WORDS)

def clean_text(text):
    text = re.sub(r"[^a-z\s]", " ", str(text).lower())
    return " ".join(w for w in text.split() if w not in STOP and len(w) > 2)

df["description_clean"] = df["description"].apply(clean_text)

# numeric features from text
df["desc_word_count"]  = df["description"].str.split().str.len().fillna(0).astype(int)
df["title_word_count"] = df["title"].str.split().str.len()
df["title_length"]     = df["title"].str.len()
df["is_series"]        = df["title"].str.contains(r"#\d+", regex=True).astype(int)   # e.g. "(Scott Pilgrim #1)"
df["has_subtitle"]     = df["title"].str.contains(":", regex=False).astype(int)

# group rare categories for modelling
MIN_CAT = 10
cat_counts = df["category"].value_counts()
df["category_group"] = df["category"].where(df["category"].map(cat_counts) >= MIN_CAT, "Other")
print("Categories:", df["category"].nunique(), "| groups used for modelling:", df["category_group"].nunique())
cat_counts.head(10)

# %% [markdown]
# ### 2.3 Outlier and consistency checks

# %%
# price outliers using the IQR rule
q1, q3 = df["price"].quantile([0.25, 0.75])
iqr = q3 - q1
low, high = q1 - 1.5 * iqr, q3 + 1.5 * iqr
n_out = int(((df["price"] < low) | (df["price"] > high)).sum())
print(f"IQR fences: {low:.2f} to {high:.2f} | prices outside: {n_out}")
print(f"Price range: {df['price'].min():.2f} to {df['price'].max():.2f}")

# consistency: price should equal price excluding tax, and tax should be constant
excl = to_number(df["price_excl_tax_raw"])
print("Rows where listing price differs from price excl. tax:", int((excl.notna() & (excl != df["price"])).sum()))
print("Distinct tax values:", df["tax_raw"].dropna().unique())
print("Distinct review counts:", df["num_reviews_raw"].dropna().unique())

# %% [markdown]
# **Decision on outliers:** I keep all prices. They are real listed prices (not typing errors) and a price range from cheap to expensive is exactly what the model needs to learn from. Removing them would hide part of the catalogue. Tax is the same for every book and the review count is constant, so those columns carry no information and are not included in the final dataset.

# %% [markdown]
# ### 2.4 Final clean dataset

# %%
final_cols = ["title", "upc", "category", "category_group", "price", "rating", "rating_word",
              "in_stock", "quantity", "desc_word_count", "title_word_count", "title_length",
              "is_series", "has_subtitle", "has_description", "description", "description_clean"]
df_clean = df[final_cols].copy()

# validation checks
assert df_clean["price"].gt(0).all(),                "non-positive price found"
assert df_clean["rating"].between(1, 5).all(),       "rating outside 1-5"
assert df_clean[["title", "price", "rating", "category", "quantity"]].notna().all().all(), "nulls in key columns"
assert df["product_url"].is_unique,                  "duplicate books remain"

df_clean.to_csv(CLEAN_PATH, index=False)
print(f"Saved {len(df_clean)} clean rows to {CLEAN_PATH}")
print(df_clean.dtypes)
df_clean.head()

# %%
df = df_clean.copy()     # use the clean data from here on
df.describe().round(2)

# %% [markdown]
# ## 3. Data Visualisation and Exploration
# I look at one question per plot. Interpretations below each chart are written from the computed numbers (so they stay accurate if the site data changes). I use Spearman correlation because rating is ordinal and price is not perfectly normal.

# %%
def save_fig(name):
    plt.tight_layout()
    plt.savefig(f"figures/{name}.png", dpi=150, bbox_inches="tight")
    plt.show()

def strength(r):
    r = abs(r)
    if np.isnan(r): return "undefined"
    if r < 0.1:  return "negligible"
    if r < 0.3:  return "weak"
    if r < 0.5:  return "moderate"
    return "strong"

insights = {}   # numbers I reuse in the report

# %% [markdown]
# ### 3.1 How are prices distributed?

# %%
fig, axes = plt.subplots(1, 2, figsize=(12, 4))
sns.histplot(df["price"], bins=30, kde=True, ax=axes[0], color="steelblue")
axes[0].axvline(df["price"].mean(), color="red", ls="--", label=f"mean £{df['price'].mean():.2f}")
axes[0].axvline(df["price"].median(), color="green", ls="--", label=f"median £{df['price'].median():.2f}")
axes[0].set_title("Distribution of book prices"); axes[0].set_xlabel("Price (£)"); axes[0].legend()
sns.boxplot(x=df["price"], ax=axes[1], color="lightsteelblue")
axes[1].set_title("Price boxplot"); axes[1].set_xlabel("Price (£)")
save_fig("fig01_price_distribution")

insights.update(price_mean=df["price"].mean(), price_median=df["price"].median(), price_std=df["price"].std(),
                price_min=df["price"].min(), price_max=df["price"].max(), price_skew=skew(df["price"]))
shape = "roughly symmetric" if abs(insights["price_skew"]) < 0.3 else ("right-skewed" if insights["price_skew"] > 0 else "left-skewed")
print(f"Prices range from £{insights['price_min']:.2f} to £{insights['price_max']:.2f}. "
      f"Mean £{insights['price_mean']:.2f}, median £{insights['price_median']:.2f}, std £{insights['price_std']:.2f}.")
print(f"Skewness = {insights['price_skew']:.2f}, so the distribution is {shape}.")

# %% [markdown]
# ### 3.2 Do ratings relate to price?

# %%
rating_stats = df.groupby("rating")["price"].agg(["count", "mean", "median"]).round(2)
print(rating_stats)

fig, axes = plt.subplots(1, 3, figsize=(15, 4))
sns.countplot(x="rating", data=df, ax=axes[0], color="steelblue")
axes[0].set_title("Number of books per star rating")
sns.boxplot(x="rating", y="price", data=df, ax=axes[1], color="lightsteelblue")
axes[1].set_title("Price by star rating"); axes[1].set_ylabel("Price (£)")
rating_stats["mean"].plot(kind="bar", ax=axes[2], color="seagreen")
axes[2].axhline(df["price"].mean(), color="red", ls="--", label="overall mean")
axes[2].set_title("Mean price by star rating"); axes[2].set_ylabel("Mean price (£)"); axes[2].legend()
axes[2].tick_params(axis="x", rotation=0)
save_fig("fig02_rating_vs_price")

rho_rating, p_rating = spearmanr(df["rating"], df["price"])
groups = [g["price"].values for _, g in df.groupby("rating")]
kw_rating_p = kruskal(*groups).pvalue if len(groups) > 1 else np.nan
insights.update(rho_rating=rho_rating, p_rating=p_rating, kw_rating_p=kw_rating_p)
print(f"\nSpearman correlation (rating vs price) = {rho_rating:.3f} (p = {p_rating:.3f}): a {strength(rho_rating)} relationship.")
print(f"Kruskal-Wallis test of price across rating groups: p = {kw_rating_p:.3f} -> "
      + ("price differs between rating groups." if kw_rating_p < 0.05 else "no statistically detectable difference in price between rating groups."))

# %% [markdown]
# ### 3.3 Genre (category) patterns

# %%
cat_stats = df.groupby("category")["price"].agg(["count", "mean", "median"]).round(2)
cat_rating = df.groupby("category")["rating"].mean().round(2)
big = cat_stats[cat_stats["count"] >= MIN_CAT].sort_values("mean")

fig, axes = plt.subplots(1, 3, figsize=(18, max(5, 0.28 * len(big))))
cat_counts.head(15).sort_values().plot(kind="barh", ax=axes[0], color="steelblue")
axes[0].set_title("15 largest categories (number of books)")
big["mean"].plot(kind="barh", ax=axes[1], color="darkorange")
axes[1].axvline(df["price"].mean(), color="red", ls="--", label="overall mean")
axes[1].set_title(f"Mean price by category (n >= {MIN_CAT})"); axes[1].set_xlabel("Mean price (£)"); axes[1].legend(loc="lower right")
cat_rating.loc[big.index].plot(kind="barh", ax=axes[2], color="seagreen")
axes[2].axvline(df["rating"].mean(), color="red", ls="--")
axes[2].set_title(f"Mean star rating by category (n >= {MIN_CAT})"); axes[2].set_xlabel("Mean rating")
save_fig("fig03_category_patterns")

big_groups = [g["price"].values for c, g in df.groupby("category") if c in big.index]
kw_cat_p = kruskal(*big_groups).pvalue if len(big_groups) > 1 else np.nan
insights.update(kw_cat_p=kw_cat_p, top_cat=big["mean"].idxmax(), top_cat_price=big["mean"].max(),
                low_cat=big["mean"].idxmin(), low_cat_price=big["mean"].min(),
                largest_cat=cat_counts.index[0], largest_cat_n=int(cat_counts.iloc[0]),
                n_categories=df["category"].nunique())
print(f"Largest category: {insights['largest_cat']} ({insights['largest_cat_n']} books).")
print(f"Highest mean price (n>={MIN_CAT}): {insights['top_cat']} at £{insights['top_cat_price']:.2f}; "
      f"lowest: {insights['low_cat']} at £{insights['low_cat_price']:.2f}.")
print(f"Kruskal-Wallis across categories: p = {kw_cat_p:.3f} -> "
      + ("category prices differ significantly." if kw_cat_p < 0.05 else "no statistically detectable price difference between categories."))

# %% [markdown]
# ### 3.4 Stock availability

# %%
fig, axes = plt.subplots(1, 2, figsize=(12, 4))
sns.histplot(df["quantity"], bins=20, ax=axes[0], color="steelblue")
axes[0].set_title("Stock quantity per book"); axes[0].set_xlabel("Units available")
sns.scatterplot(x="quantity", y="price", data=df, alpha=0.4, ax=axes[1])
axes[1].set_title("Price vs stock quantity"); axes[1].set_ylabel("Price (£)")
save_fig("fig04_stock")

insights["pct_in_stock"] = df["in_stock"].mean() * 100
if df["quantity"].nunique() > 1:
    rho_qty, p_qty = spearmanr(df["quantity"], df["price"])
else:
    rho_qty, p_qty = np.nan, np.nan
insights.update(rho_qty=rho_qty, p_qty=p_qty, qty_mean=df["quantity"].mean())
print(f"{insights['pct_in_stock']:.1f}% of books are in stock; average quantity {insights['qty_mean']:.1f} units.")
print(f"Spearman correlation (quantity vs price) = {rho_qty:.3f} (p = {p_qty:.3f}): {strength(rho_qty)}.")

# %% [markdown]
# ### 3.5 Title and description features

# %%
fig, axes = plt.subplots(1, 3, figsize=(16, 4))
d = df[df["has_description"] == 1]
sns.regplot(x="desc_word_count", y="price", data=d, scatter_kws={"alpha": 0.3}, line_kws={"color": "red"}, ax=axes[0])
axes[0].set_title("Price vs description length"); axes[0].set_xlabel("Description words")
sns.regplot(x="title_word_count", y="price", data=df, x_jitter=0.2, scatter_kws={"alpha": 0.3}, line_kws={"color": "red"}, ax=axes[1])
axes[1].set_title("Price vs title length"); axes[1].set_xlabel("Title words")
sns.boxplot(x="is_series", y="price", data=df, ax=axes[2], color="lightsteelblue")
axes[2].set_title("Price: standalone (0) vs series (1)")
save_fig("fig05_text_features")

rho_desc, _ = spearmanr(d["desc_word_count"], d["price"])
rho_title, _ = spearmanr(df["title_word_count"], df["price"])
series_gap = df.groupby("is_series")["price"].mean()
insights.update(rho_desc=rho_desc, rho_title=rho_title, n_series=int(df["is_series"].sum()),
                series_gap=(series_gap.get(1, np.nan) - series_gap.get(0, np.nan)))
print(f"Description length vs price: rho = {rho_desc:.3f} ({strength(rho_desc)}).")
print(f"Title length vs price: rho = {rho_title:.3f} ({strength(rho_title)}).")
print(f"Series titles: {insights['n_series']} books; mean price gap (series - standalone) = £{insights['series_gap']:.2f}.")

# %% [markdown]
# ### 3.6 Correlation overview

# %%
num_cols = [c for c in ["price", "rating", "quantity", "desc_word_count", "title_word_count",
                        "title_length", "is_series", "has_subtitle"] if df[c].nunique() > 1]
corr = df[num_cols].corr(method="spearman")
plt.figure(figsize=(8, 6))
sns.heatmap(corr, annot=True, fmt=".2f", cmap="coolwarm", center=0, vmin=-1, vmax=1)
plt.title("Spearman correlation matrix")
save_fig("fig06_correlation")

price_corr = corr["price"].drop("price").sort_values(key=abs, ascending=False)
insights["strongest_feature"] = price_corr.index[0]
insights["strongest_corr"] = price_corr.iloc[0]
print("Correlation of each feature with price:")
print(price_corr.round(3))
print(f"\nTitle word count and title length correlate at {corr.loc['title_word_count','title_length']:.2f} "
      "(so I will only keep one of them in the model to avoid multicollinearity).")

# %% [markdown]
# ### 3.7 What words appear in descriptions? (keyword 'tags')

# %%
def top_keywords(texts, n=10):
    vec = CountVectorizer(binary=True, min_df=3)
    X = vec.fit_transform(texts)
    freq = pd.Series(np.asarray(X.sum(axis=0)).ravel() / X.shape[0], index=vec.get_feature_names_out())
    return freq.sort_values(ascending=False).head(n)

d = df[df["has_description"] == 1]
q_low, q_high = d["price"].quantile([0.25, 0.75])
cheap_kw     = top_keywords(d.loc[d["price"] <= q_low,  "description_clean"])
expensive_kw = top_keywords(d.loc[d["price"] >= q_high, "description_clean"])
overall_kw   = top_keywords(d["description_clean"], 15)

fig, axes = plt.subplots(1, 3, figsize=(16, 4))
overall_kw.sort_values().plot(kind="barh", ax=axes[0], color="steelblue"); axes[0].set_title("Top keywords (all books)")
cheap_kw.sort_values().plot(kind="barh", ax=axes[1], color="seagreen");    axes[1].set_title("Cheapest quartile")
expensive_kw.sort_values().plot(kind="barh", ax=axes[2], color="darkorange"); axes[2].set_title("Most expensive quartile")
for ax in axes: ax.set_xlabel("Share of descriptions containing the word")
save_fig("fig07_keywords")

insights["top_keywords"] = list(overall_kw.index[:5])
print("Most common description keywords:", ", ".join(insights["top_keywords"]))
print("Cheapest quartile keywords:", ", ".join(cheap_kw.index[:6]))
print("Most expensive quartile keywords:", ", ".join(expensive_kw.index[:6]))

# %% [markdown]
# ### 3.8 Price bands across ratings

# %%
bands = pd.cut(df["price"], bins=[0, 20, 30, 40, 50, 1000], labels=["<£20", "£20-30", "£30-40", "£40-50", ">£50"])
share = pd.crosstab(bands, df["rating"], normalize="columns")
plt.figure(figsize=(8, 4.5))
sns.heatmap(share, annot=True, fmt=".0%", cmap="Blues")
plt.title("Share of each rating group falling in each price band")
plt.xlabel("Star rating"); plt.ylabel("Price band")
save_fig("fig08_priceband_rating")
print("If each column looks similar, price band does not depend much on star rating.")

# %% [markdown]
# ## 4. Predictive Modelling (Linear Regression)
#
# ### 4.1 Features and justification
# **Target:** `price` (£).
#
# | Feature | Why include it |
# |---|---|
# | `rating` (1 to 5) | main quality signal customers see; ordinal so used as a number |
# | `quantity` | stock level may reflect demand or supply strategy |
# | `desc_word_count`, `has_description` | richer listings may be priced differently; the flag handles books with no description |
# | `title_word_count` | simple listing property; I drop `title_length` because it is almost the same information (high correlation, multicollinearity) |
# | `is_series`, `has_subtitle` | series books and subtitled books may follow different pricing |
# | `category_group` (one-hot) | genre is the main business-level descriptor |
#
# `in_stock` is left out automatically if every book is in stock (a constant column gives the model nothing). I do **not** scale features: ordinary Linear Regression is unaffected by scaling and keeping the original units makes the coefficients easy to read in pounds.

# %%
base_features = ["rating", "quantity", "desc_word_count", "has_description",
                 "title_word_count", "is_series", "has_subtitle", "in_stock"]
base_features = [f for f in base_features if df[f].nunique() > 1]       # drop constants
print("Numeric features used:", base_features)

# one-hot encode category_group; use the biggest group as the reference level
reference = df["category_group"].value_counts().index[0]
X_struct = df[base_features].copy()
X_full = pd.concat([X_struct, pd.get_dummies(df["category_group"], prefix="cat", dtype=int)
                    .drop(columns=f"cat_{reference}")], axis=1)
y = df["price"]
print(f"Reference category (baseline for the dummies): {reference}")
print("Modelling matrix:", X_full.shape, "| missing values:", int(X_full.isna().sum().sum()))

# %% [markdown]
# ### 4.2 Train/test split and model fitting
# I hold out 20% of the books as a test set the model never sees, so the test metrics show how well it generalises. A baseline that always predicts the mean price shows whether the model learns anything.

# %%
X_train, X_test, y_train, y_test = train_test_split(X_full, y, test_size=0.2, random_state=42)
print("Train:", X_train.shape, "| Test:", X_test.shape)

model = LinearRegression().fit(X_train, y_train)
baseline = DummyRegressor(strategy="mean").fit(X_train, y_train)

def evaluate(y_true, y_pred):
    return {"MAE (£)": mean_absolute_error(y_true, y_pred),
            "RMSE (£)": np.sqrt(mean_squared_error(y_true, y_pred)),
            "R2": r2_score(y_true, y_pred)}

results = pd.DataFrame({
    "Linear Regression (train)": evaluate(y_train, model.predict(X_train)),
    "Linear Regression (test)":  evaluate(y_test,  model.predict(X_test)),
    "Mean baseline (test)":      evaluate(y_test,  baseline.predict(X_test)),
}).T.round(3)
results

# %% [markdown]
# ### 4.3 Cross-validation and feature-set comparison

# %%
kf = KFold(n_splits=5, shuffle=True, random_state=42)
cv_full   = cross_val_score(LinearRegression(), X_full,   y, cv=kf, scoring="r2")
cv_struct = cross_val_score(LinearRegression(), X_struct, y, cv=kf, scoring="r2")
cv_table = pd.DataFrame({"Mean CV R2": [cv_struct.mean(), cv_full.mean()],
                         "Std": [cv_struct.std(), cv_full.std()]},
                        index=["Without category", "With category (final model)"]).round(3)
print(cv_table)

metrics = results.loc["Linear Regression (test)"]
insights.update(r2_test=metrics["R2"], mae_test=metrics["MAE (£)"], rmse_test=metrics["RMSE (£)"],
                r2_train=results.loc["Linear Regression (train)", "R2"],
                mae_base=results.loc["Mean baseline (test)", "MAE (£)"],
                cv_mean=cv_full.mean(), cv_std=cv_full.std(), cv_struct_mean=cv_struct.mean(),
                n_train=len(X_train), n_test=len(X_test), n_features=X_full.shape[1])
print(f"\nTest R2 = {insights['r2_test']:.3f}, MAE = £{insights['mae_test']:.2f} (baseline MAE £{insights['mae_base']:.2f}).")

# %% [markdown]
# ### 4.4 Diagnostics and coefficients

# %%
pred_test = model.predict(X_test)
resid = y_test - pred_test

fig, axes = plt.subplots(1, 3, figsize=(16, 4))
axes[0].scatter(y_test, pred_test, alpha=0.5)
lims = [min(y_test.min(), pred_test.min()), max(y_test.max(), pred_test.max())]
axes[0].plot(lims, lims, "r--"); axes[0].set_xlabel("Actual price (£)"); axes[0].set_ylabel("Predicted price (£)")
axes[0].set_title("Actual vs predicted (test set)")
axes[1].scatter(pred_test, resid, alpha=0.5); axes[1].axhline(0, color="r", ls="--")
axes[1].set_xlabel("Predicted price (£)"); axes[1].set_ylabel("Residual (£)"); axes[1].set_title("Residuals vs predicted")
sns.histplot(resid, bins=20, kde=True, ax=axes[2], color="steelblue"); axes[2].set_title("Residual distribution")
save_fig("fig09_model_diagnostics")

coef = pd.Series(model.coef_, index=X_full.columns).sort_values(key=abs, ascending=False)
top_coef = coef.head(12).sort_values()
plt.figure(figsize=(8, 5))
top_coef.plot(kind="barh", color=["seagreen" if v > 0 else "indianred" for v in top_coef])
plt.axvline(0, color="black", lw=0.8); plt.xlabel("Coefficient (change in price, £)")
plt.title("Largest Linear Regression coefficients")
save_fig("fig10_coefficients")

insights["resid_std"] = resid.std()
insights["resid_skew"] = skew(resid)
insights["intercept"] = model.intercept_
print(f"Intercept: £{model.intercept_:.2f}")
print(coef.round(3).head(12))
if "rating" in coef.index:
    print(f"\nHolding other features fixed, one extra star changes the predicted price by {coef['rating']:+.2f} pounds.")
    insights["rating_coef"] = coef["rating"]

# %% [markdown]
# ### 4.5 Interpretation of results and limitations
# The cell below writes the interpretation from the actual numbers, so the text always matches the output.

# %%
r2, mae, base = insights["r2_test"], insights["mae_test"], insights["mae_base"]
if r2 < 0.10:
    level = "very little of the variation in price"
elif r2 < 0.40:
    level = "a modest share of the variation in price"
else:
    level = "a substantial share of the variation in price"

print(f"1. The model explains {level} (test R2 = {r2:.2f}; cross-validated R2 = {insights['cv_mean']:.2f} +/- {insights['cv_std']:.2f}).")
print(f"2. Typical error is about £{mae:.2f} (MAE) against £{base:.2f} for just guessing the mean"
      + (", so it clearly beats the baseline." if (mae < 0.95 * base and r2 > 0) else ", so it is no better than simply guessing the mean price."))
print(f"3. Train R2 = {insights['r2_train']:.2f} vs test R2 = {r2:.2f}: "
      + ("a noticeable gap suggests some overfitting from the many category dummies." if insights["r2_train"] - r2 > 0.10
         else "a small gap, so overfitting is not a major concern."))
print(f"4. Adding category changes mean CV R2 from {insights['cv_struct_mean']:.2f} to {insights['cv_mean']:.2f}.")
print("\nLimitations:")
print("- Only public listing fields are available: no author, publisher, page count, publication year or sales data.")
print("- Linear Regression assumes straight-line, additive effects; real pricing may be non-linear.")
print("- The site states that prices and ratings are randomly assigned, so weak predictability is expected here.")
print(f"- Dataset is modest in size ({len(df)} books) and some categories have few books, so individual category coefficients are noisy.")
print("- Residuals should be roughly centred on zero with constant spread; check the diagnostic plots for patterns.")

# %% [markdown]
# ## 5. Professional Report (auto-generated from the results above)
# This cell assembles a consulting-style report using the numbers computed in this notebook and saves it to `report/Project_Report.docx` (and a `.md` copy). Because every number comes from the variables above, the report always matches the notebook. Add your name and read it through before submitting.

# %%
def fmt_p(p):
    return "p < 0.001" if p < 0.001 else f"p = {p:.3f}"

I = insights
r2_label = "very low" if I["r2_test"] < 0.10 else ("moderate" if I["r2_test"] < 0.40 else "fairly strong")
rating_sig = I["p_rating"] < 0.05
cat_sig = I["kw_cat_p"] < 0.05
weak_model = I["r2_test"] < 0.30

# titles priced lowest relative to model expectation (full-data predictions, for illustration only)
df_pred = df.copy()
df_pred["predicted"] = model.predict(X_full)
df_pred["gap"] = df_pred["price"] - df_pred["predicted"]
cheap_titles = df_pred.nsmallest(5, "gap")[["title", "category", "price", "predicted"]].round(2)

R = []   # report items: (kind, content)
def add(kind, content=None, extra=None): R.append((kind, content, extra))

add("title", "Book Catalogue Analysis and Price Prediction: Books to Scrape")
add("p", "Prepared for: Publishing Analytics Team  |  Prepared by: ____________________")

add("h1", "1. Executive Summary")
add("p", f"We scraped {len(df)} books across {I['n_categories']} categories from Books to Scrape, cleaned the data, explored price, rating, genre and stock patterns, and fitted a Linear Regression to estimate price. "
         f"Prices average £{I['price_mean']:.2f} (median £{I['price_median']:.2f}) and range from £{I['price_min']:.2f} to £{I['price_max']:.2f}. "
         f"The final model explains {r2_label} variation in price (test R² = {I['r2_test']:.2f}) with a typical error of £{I['mae_test']:.2f}.")
add("bullet", f"Star rating and price have a {strength(I['rho_rating'])} relationship (Spearman ρ = {I['rho_rating']:.2f}, {fmt_p(I['p_rating'])}).")
add("bullet", (f"Category prices differ significantly ({fmt_p(I['kw_cat_p'])}); {I['top_cat']} is highest (£{I['top_cat_price']:.2f}) and {I['low_cat']} lowest (£{I['low_cat_price']:.2f})." if cat_sig
               else f"Category does not produce statistically detectable price differences ({fmt_p(I['kw_cat_p'])}); category means range from £{I['low_cat_price']:.2f} ({I['low_cat']}) to £{I['top_cat_price']:.2f} ({I['top_cat']})."))
add("bullet", ("The available public fields predict price only weakly, so a pricing tool built on them alone would not be reliable." if weak_model
               else "The available public fields give useful, though imperfect, price estimates."))

add("h1", "2. Methodology")
add("h2", "2.1 Data acquisition")
add("p", f"Data was collected with Python (requests and BeautifulSoup). The scraper walks all listing pages by following the next-page link, reads title, price, star rating (a CSS class) and availability, and then opens every product page for category, exact stock quantity and description. Requests are delayed, retried on failure and use a User-Agent header. The raw dataset has {len(raw)} rows" + (", above the 200-book minimum." if len(raw) >= 200 else ", which is BELOW the 200-book minimum; re-run the scrape.") )
add("h2", "2.2 Data cleaning")
add("p", "Prices were converted from text (e.g. '£51.77') to numbers; star-rating words (One to Five) were mapped to 1 to 5; availability text was split into an in-stock flag and a quantity; descriptions were stripped of scraping artefacts and cleaned (lower-case, stop-words removed). Books without a description were kept and flagged rather than deleted. Rows with no price or rating would have been dropped because price is the target. Categories with fewer than 10 books were grouped as 'Other' for modelling. Constant columns (tax, product type, review count) were removed. Price outliers were reviewed and retained because they are valid listed prices.")
add("h2", "2.3 Exploration and modelling")
add("p", f"Relationships were explored with histograms, box plots, bar charts, scatter plots, a correlation matrix and keyword frequencies. The model is an ordinary least squares Linear Regression predicting price from rating, stock quantity, description length, title length, series and subtitle indicators and category. Data was split 80/20 into training ({I['n_train']}) and test ({I['n_test']}) sets, evaluated with MAE, RMSE and R², compared with a mean-price baseline, and checked with 5-fold cross-validation.")

add("h1", "3. Key Insights")
add("h2", "3.1 Pricing landscape")
add("p", f"Price is {'roughly symmetric' if abs(I['price_skew']) < 0.3 else ('right-skewed' if I['price_skew'] > 0 else 'left-skewed')} (skewness {I['price_skew']:.2f}) with a standard deviation of £{I['price_std']:.2f}, so prices are spread widely across the range rather than clustered around one level.")
add("img", "figures/fig01_price_distribution.png", "Figure 1: Distribution of book prices")
add("h2", "3.2 Ratings")
add("p", f"Mean price by star rating is shown below. The relationship between rating and price is {strength(I['rho_rating'])} (ρ = {I['rho_rating']:.2f}); "
         + ("price does differ between rating groups." if rating_sig else "price does not differ detectably between rating groups, so a higher rating does not imply a higher price in this catalogue."))
add("img", "figures/fig02_rating_vs_price.png", "Figure 2: Rating counts, price by rating and mean price by rating")
add("h2", "3.3 Genres")
add("p", f"The largest category is {I['largest_cat']} with {I['largest_cat_n']} books. " +
         (f"Genre matters for price: {I['top_cat']} averages £{I['top_cat_price']:.2f} versus £{I['low_cat_price']:.2f} for {I['low_cat']}." if cat_sig
          else f"Differences between genres are small relative to the spread within each genre ({fmt_p(I['kw_cat_p'])})."))
add("img", "figures/fig03_category_patterns.png", "Figure 3: Category sizes, mean price and mean rating")
add("h2", "3.4 Stock, titles and descriptions")
add("p", f"{I['pct_in_stock']:.0f}% of books are in stock, with an average of {I['qty_mean']:.1f} units. Correlation of stock quantity with price is {strength(I['rho_qty'])} (ρ = {I['rho_qty']:.2f}). "
         f"Description length (ρ = {I['rho_desc']:.2f}) and title length (ρ = {I['rho_title']:.2f}) show {strength(max(abs(I['rho_desc']), abs(I['rho_title'])))} relationships with price. "
         f"Series titles ({I['n_series']} books) differ from standalone titles by £{I['series_gap']:.2f} on average. Common description words include {', '.join(I['top_keywords'])}.")
add("img", "figures/fig05_text_features.png", "Figure 4: Price against description length, title length and series status")
add("img", "figures/fig06_correlation.png", "Figure 5: Correlation matrix")

add("h1", "4. Modelling Outcomes")
mt = results.reset_index().rename(columns={"index": "Model"}).round(3)
add("table", mt, "Table 1: Linear Regression performance")
add("p", f"On unseen test data the model achieves R² = {I['r2_test']:.2f}, MAE = £{I['mae_test']:.2f} and RMSE = £{I['rmse_test']:.2f}, compared with a baseline MAE of £{I['mae_base']:.2f} for always predicting the mean ("
         + ("the model improves on the baseline" if (I['mae_test'] < 0.95 * I['mae_base'] and I['r2_test'] > 0) else "the model is not meaningfully better than the baseline") + "). "
         f"Five-fold cross-validated R² is {I['cv_mean']:.2f} (± {I['cv_std']:.2f}); without category it is {I['cv_struct_mean']:.2f}. "
         + ("Because R² is low, most of the variation in price is not captured by these features." if weak_model else "The model captures a meaningful part of price variation.")
         + (" A negative test R² means the model predicted unseen books worse than simply using the average price." if I["r2_test"] < 0 else ""))
ct = coef.head(8).round(3).reset_index(); ct.columns = ["Feature", "Coefficient (£)"]
add("table", ct, "Table 2: Largest coefficients")
add("img", "figures/fig10_coefficients.png", "Figure 6: Largest regression coefficients")
add("img", "figures/fig09_model_diagnostics.png", "Figure 7: Model diagnostics")
add("p", "Limitations: only public listing fields were available (no author, publisher, page count, publication year or sales); Linear Regression assumes additive, straight-line effects; the dataset is modest in size with many small categories; and the source site states that its prices and ratings were randomly assigned, so weak predictability is expected and conclusions should not be generalised to real bookstores.")

add("h1", "5. Practical Implications")
add("bullet", ("Star rating alone should not be used to set or judge price in this catalogue." if abs(I["rho_rating"]) < 0.3 else "Star rating carries some pricing information, but it should be combined with other signals."))
add("bullet", ("Genre-specific price benchmarks are justified by the observed category differences." if cat_sig else "Separate genre price tiers are not supported by this data; genre should not be the main basis for pricing decisions."))
add("bullet", ("Price cannot be reliably predicted from listing fields alone; richer data is needed before automating pricing." if weak_model else "Listing fields can support a first-pass price estimate that analysts then review."))

add("h1", "6. Recommendations")
add("h2", "Retailers")
add("bullet", "Benchmark each title against the median price of its category (Figure 3) rather than against its star rating.")
add("bullet", "Use the model only as a sanity check for outlier prices, not as an automatic pricing engine, unless performance improves.")
add("h2", "Publishers")
add("bullet", "Collect and test additional drivers such as page count, format, author profile, publication year and sales history before setting list prices.")
add("bullet", "Track whether series titles justify a premium or discount in your own sales data." )
add("h2", "Data-driven book buyers")
add("bullet", "Do not assume a higher-rated book costs more or less; compare price within the same category.")
add("bullet", "Titles priced furthest below the model's expectation may be worth a closer look, treating the list as a starting point only because model accuracy is limited (Table 3).")
add("table", cheap_titles.reset_index(drop=True), "Table 3: Titles priced furthest below model expectation")

add("h1", "7. Next Steps")
add("p", "Add richer features (publisher, page count, publication year, author), test non-linear models such as random forests for comparison, and repeat the analysis on a dataset where prices are real.")

# ---------- render to Markdown ----------
def df_to_md(t):
    cols = list(t.columns)
    lines = ["| " + " | ".join(map(str, cols)) + " |", "|" + "---|" * len(cols)]
    for _, row in t.iterrows():
        lines.append("| " + " | ".join(str(v) for v in row) + " |")
    return "\n".join(lines)

md_lines = []
for kind, content, extra in R:
    if kind == "title":  md_lines.append(f"# {content}\n")
    elif kind == "h1":   md_lines.append(f"\n## {content}\n")
    elif kind == "h2":   md_lines.append(f"\n### {content}\n")
    elif kind == "p":    md_lines.append(f"{content}\n")
    elif kind == "bullet": md_lines.append(f"- {content}")
    elif kind == "img":  md_lines.append(f"\n![{extra}](../{content})\n*{extra}*\n")
    elif kind == "table": md_lines.append(f"\n*{extra}*\n\n{df_to_md(content)}\n")
with open("report/Project_Report.md", "w", encoding="utf-8") as f:
    f.write("\n".join(md_lines))
print("Saved report/Project_Report.md")

# ---------- render to Word ----------
try:
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    for kind, content, extra in R:
        if kind == "title":    doc.add_heading(content, 0)
        elif kind == "h1":     doc.add_heading(content, 1)
        elif kind == "h2":     doc.add_heading(content, 2)
        elif kind == "p":      doc.add_paragraph(content)
        elif kind == "bullet": doc.add_paragraph(content, style="List Bullet")
        elif kind == "img":
            doc.add_picture(content, width=Inches(6.0))
            doc.add_paragraph().add_run(extra).italic = True
        elif kind == "table":
            doc.add_paragraph(extra)
            t = doc.add_table(rows=1, cols=len(content.columns)); t.style = "Light Grid Accent 1"
            for j, c in enumerate(content.columns): t.rows[0].cells[j].text = str(c)
            for _, row in content.iterrows():
                cells = t.add_row().cells
                for j, v in enumerate(row): cells[j].text = str(v)
    doc.save("report/Project_Report.docx")
    print("Saved report/Project_Report.docx")
except ImportError:
    print("python-docx not installed (pip install python-docx); the Markdown report was still saved.")

# %% [markdown]
# ## 6. Deliverables checklist
# | Deliverable | File |
# |---|---|
# | Jupyter Notebook | this notebook |
# | Final dataset (clean) | `data/books_clean.csv` |
# | Raw scraped data | `data/books_raw.csv` |
# | Report | `report/Project_Report.docx` (and `.md`) |
# | Figures | `figures/` |
